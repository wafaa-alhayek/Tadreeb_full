"""FR-03: import students from an Excel sheet, one student per row."""

import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models.functions import Lower
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from accounts.models import Department, Specialization, StudentProfile, User

MAX_FILE_BYTES = 5 * 1024 * 1024  # Same limit as uploads (BR-08).
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# (header in the sheet, field, required)
COLUMNS = [
    ("الرقم الجامعي", "username", True),
    ("الاسم الأول", "first_name", True),
    ("اسم العائلة", "last_name", True),
    ("البريد الإلكتروني", "email", True),
    ("القسم", "department", True),
    ("التخصص", "specialization", True),
    ("الساعات المنجزة", "completed_credit_hours", True),
    ("حالة القيد", "academic_status", True),
    ("المعدل", "gpa", False),
    ("الهاتف", "phone", False),
]
HEADERS = {field: header for header, field, _ in COLUMNS}
EXAMPLE_ROW = [
    "120260001", "أحمد", "خالد", "ahmad@example.com", "تكنولوجيا المعلومات",
    "برمجة", 95, "منتظم", 3.2, "0599123456",
]

# The sheet uses the Arabic names (منتظم / موقوف / متخرج); the codes are accepted too.
STATUS_BY_NAME = {label: value for value, label in StudentProfile.AcademicStatus.choices}
STATUS_BY_NAME.update({value: value for value in StudentProfile.AcademicStatus.values})

# Re-importing an existing student updates these (a new sheet comes every semester).
PROFILE_FIELDS = [
    "department", "specialization", "completed_credit_hours", "gpa", "academic_status", "phone",
]


def build_workbook(rows):
    """An .xlsx file (bytes) with the header row and the given rows, right to left (NFR-11)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الطلاب"
    sheet.sheet_view.rightToLeft = True
    sheet.append([header for header, _, _ in COLUMNS])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index in range(1, len(COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 20
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_template():
    return build_workbook([EXAMPLE_ROW])


def import_students(file, dry_run=False):
    """Check every row; save all of them only if none has an error."""
    raw_rows, errors = read_sheet(file)
    rows = []
    if not errors:
        rows, errors = check_rows(raw_rows)

    created = sum(1 for row in rows if row["user"] is None)
    saved = not errors and not dry_run
    if saved:
        save_students(rows)
    return {
        "dry_run": dry_run,
        "saved": saved,
        "created": created,
        "updated": len(rows) - created,
        "errors": errors,
    }


def row_error(row, field, message):
    return {"row": row, "column": HEADERS.get(field), "message": message}


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 120260001 as 120260001.0
    return str(value).strip()


def read_sheet(file):
    """Return ([(row number, {field: text})], errors) from the first sheet."""
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        # A renamed or damaged file fails in many different ways inside openpyxl.
        return [], [row_error(None, None, "تعذّرت قراءة الملف. تأكد أنه ملف Excel (.xlsx) سليم.")]

    sheet_rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = next(sheet_rows, None) or ()
    positions = {as_text(title): index for index, title in enumerate(header)}
    missing = [h for h, _, required in COLUMNS if required and h not in positions]
    if missing:
        workbook.close()
        return [], [row_error(1, None, "أعمدة ناقصة في الملف: " + "، ".join(missing))]

    raw_rows = []
    for number, values in enumerate(sheet_rows, start=2):
        raw = {}
        for header_text, field, _ in COLUMNS:
            index = positions.get(header_text)
            raw[field] = as_text(values[index]) if index is not None and index < len(values) else ""
        if any(raw.values()):  # Skip empty lines.
            raw_rows.append((number, raw))
    workbook.close()

    if not raw_rows:
        return [], [row_error(None, None, "الملف لا يحتوي على أي طالب.")]
    return raw_rows, []


def check_rows(raw_rows):
    """Turn each row into clean values, or collect its errors."""
    departments = {d.name: d for d in Department.objects.all()}
    specializations = {(s.department_id, s.name): s for s in Specialization.objects.all()}
    specialization_names = {name for _, name in specializations}

    usernames = {raw["username"] for _, raw in raw_rows}
    emails = {raw["email"].lower() for _, raw in raw_rows}
    existing_users = {
        user.username: user
        for user in User.objects.filter(username__in=usernames).select_related("student_profile")
    }
    email_owners = dict(
        User.objects.annotate(email_lower=Lower("email"))
        .filter(email_lower__in=emails)
        .values_list("email_lower", "username")
    )

    rows, errors = [], []
    username_rows, email_rows = {}, {}
    for number, raw in raw_rows:
        row_errors = [
            row_error(number, field, "هذا الحقل مطلوب.")
            for _, field, required in COLUMNS
            if required and not raw[field]
        ]
        username, email = raw["username"], raw["email"].lower()

        user = existing_users.get(username)
        if username:
            if username in username_rows:
                row_errors.append(row_error(
                    number, "username", f"الرقم الجامعي مكرر في الملف (السطر {username_rows[username]})."
                ))
            else:
                username_rows[username] = number
            try:
                User.username_validator(username)
            except ValidationError:
                row_errors.append(row_error(number, "username", "الرقم الجامعي يحتوي رموزًا غير مسموحة."))
            if user is not None and user.role != User.Role.STUDENT:
                row_errors.append(row_error(
                    number, "username", "يوجد مستخدم بهذا الرقم وليس طالبًا، فلا يمكن تعديله من الاستيراد."
                ))

        if email:
            if email in email_rows:
                row_errors.append(row_error(
                    number, "email", f"البريد الإلكتروني مكرر في الملف (السطر {email_rows[email]})."
                ))
            else:
                email_rows[email] = number
            try:
                validate_email(email)
            except ValidationError:
                row_errors.append(row_error(number, "email", "البريد الإلكتروني غير صحيح."))
            owner = email_owners.get(email)
            if owner is not None and owner != username:
                row_errors.append(row_error(number, "email", "البريد الإلكتروني مستخدم لحساب آخر."))

        for field, limit in (("username", 150), ("first_name", 150), ("last_name", 150), ("phone", 20)):
            if len(raw[field]) > limit:
                row_errors.append(row_error(number, field, f"النص أطول من {limit} حرفًا."))

        department = departments.get(raw["department"])
        if raw["department"] and department is None:
            row_errors.append(row_error(number, "department", f"القسم «{raw['department']}» غير موجود."))

        specialization = None
        if raw["specialization"] and department is not None:
            specialization = specializations.get((department.id, raw["specialization"]))
            if specialization is None:
                if raw["specialization"] in specialization_names:
                    message = f"التخصص «{raw['specialization']}» لا يتبع قسم «{department.name}»."
                else:
                    message = f"التخصص «{raw['specialization']}» غير موجود."
                row_errors.append(row_error(number, "specialization", message))

        hours = None
        if raw["completed_credit_hours"]:
            if re.fullmatch(r"\d+", raw["completed_credit_hours"]):
                hours = int(raw["completed_credit_hours"])
            else:
                row_errors.append(row_error(
                    number, "completed_credit_hours", "الساعات المنجزة لازم تكون رقمًا صحيحًا 0 أو أكثر."
                ))

        status = STATUS_BY_NAME.get(raw["academic_status"])
        if raw["academic_status"] and status is None:
            row_errors.append(row_error(
                number, "academic_status", "حالة القيد لازم تكون: منتظم أو موقوف أو متخرج."
            ))

        gpa = None
        if raw["gpa"]:
            try:
                gpa = Decimal(raw["gpa"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            except InvalidOperation:
                gpa = None
            if gpa is None or not Decimal(0) <= gpa <= Decimal(4):
                row_errors.append(row_error(number, "gpa", "المعدل لازم يكون رقمًا بين 0 و 4."))

        if row_errors:
            errors.extend(row_errors)
            continue
        rows.append({
            "user": user,
            "username": username,
            "first_name": raw["first_name"],
            "last_name": raw["last_name"],
            "email": email,
            "department": department,
            "specialization": specialization,
            "completed_credit_hours": hours,
            "gpa": gpa,
            "academic_status": status,
            "phone": raw["phone"],
        })
    return rows, errors


@transaction.atomic
def save_students(rows):
    """All rows or none: if one save fails, the whole import is rolled back."""
    for row in rows:
        user = row["user"]
        if user is None:
            user = User(
                username=row["username"],
                first_name=row["first_name"],
                last_name=row["last_name"],
                email=row["email"],
                role=User.Role.STUDENT,
            )
            # The first password is still an open team decision; until then the
            # account exists but nobody can log in with it (NFR-01: no plain-text passwords).
            user.set_unusable_password()
            user.save()
            profile = StudentProfile(user=user)
        else:
            # Existing student: update the academic record, never the password.
            user.email = row["email"]
            user.save(update_fields=["email"])
            profile = getattr(user, "student_profile", None) or StudentProfile(user=user)

        for field in PROFILE_FIELDS:
            setattr(profile, field, row[field])
        profile.save()

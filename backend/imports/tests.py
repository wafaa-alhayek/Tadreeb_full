import shutil
import tempfile
from io import BytesIO, StringIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import CommandError, call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.eligibility import check_eligibility
from accounts.models import Department, Specialization, StudentProfile, User
from applications.models import Application, ExceptionRequest, Placement
from applications.status import Status
from certificates.models import Certificate
from notifications.models import Notification
from providers.models import Opportunity, Provider
from training.models import DailyLog, Report

from .students import COLUMNS, build_workbook

URL = "/api/imports/students/"
TEMPLATE_URL = "/api/imports/students/template/"


def student_row(username="120260001", email=None, **changes):
    row = {
        "username": username,
        "first_name": "أحمد",
        "last_name": "خالد",
        "email": email or f"{username}@example.com",
        "department": "تكنولوجيا المعلومات",
        "specialization": "برمجة",
        "completed_credit_hours": 95,
        "academic_status": "منتظم",
        "gpa": 3.25,
        "phone": "0599123456",
    }
    row.update(changes)
    return [row[field] for _, field, _ in COLUMNS]


def excel_file(rows, name="students.xlsx"):
    return SimpleUploadedFile(name, build_workbook(rows))


class StudentImportTests(APITestCase):
    def setUp(self):
        self.coordinator = User.objects.create_user(
            username="coordinator", password="Test12345!", role=User.Role.COORDINATOR
        )
        it = Department.objects.create(name="تكنولوجيا المعلومات")
        business = Department.objects.create(name="العلوم الإدارية والمالية")
        self.programming = Specialization.objects.create(department=it, name="برمجة")
        Specialization.objects.create(department=it, name="شبكات")
        Specialization.objects.create(department=business, name="محاسبة")
        self.client.force_authenticate(self.coordinator)

    def upload(self, rows, query=""):
        return self.client.post(URL + query, {"file": excel_file(rows)}, format="multipart")

    def assert_rejected(self, response, column, text):
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["saved"])
        messages = [e["message"] for e in response.data["errors"] if e["column"] == column]
        self.assertTrue(any(text in m for m in messages), response.data["errors"])

    def test_valid_file_creates_students(self):
        response = self.upload([student_row("120260001"), student_row("120260002")])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["created"], 2)
        self.assertTrue(response.data["saved"])

        user = User.objects.get(username="120260001")
        self.assertEqual(user.role, User.Role.STUDENT)
        self.assertEqual(user.get_full_name(), "أحمد خالد")
        self.assertFalse(user.has_usable_password())  # No plain or default password (NFR-01).
        profile = user.student_profile
        self.assertEqual(profile.specialization, self.programming)
        self.assertEqual(profile.completed_credit_hours, 95)
        self.assertEqual(str(profile.gpa), "3.25")
        self.assertEqual(profile.academic_status, StudentProfile.AcademicStatus.ACTIVE)

    def test_arabic_status_names_and_optional_columns(self):
        response = self.upload([student_row(academic_status="موقوف", gpa=None, phone=None)])
        self.assertEqual(response.status_code, 200, response.data)
        profile = StudentProfile.objects.get(user__username="120260001")
        self.assertEqual(profile.academic_status, StudentProfile.AcademicStatus.SUSPENDED)
        self.assertIsNone(profile.gpa)

    def test_dry_run_saves_nothing(self):
        response = self.upload([student_row()], query="?dry_run=1")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["created"], 1)
        self.assertFalse(response.data["saved"])
        self.assertFalse(User.objects.filter(username="120260001").exists())

    def test_one_bad_row_cancels_the_whole_file(self):
        response = self.upload([student_row("120260001"), student_row("120260002", academic_status="غائب")])
        self.assert_rejected(response, "حالة القيد", "منتظم")
        self.assertEqual(response.data["errors"][0]["row"], 3)
        self.assertFalse(User.objects.filter(username="120260001").exists())

    def test_duplicate_university_id_in_file(self):
        response = self.upload([student_row("120260001"), student_row("120260001", email="b@example.com")])
        self.assert_rejected(response, "الرقم الجامعي", "مكرر")

    def test_duplicate_email_in_file(self):
        response = self.upload([
            student_row("120260001", email="same@example.com"),
            student_row("120260002", email="SAME@example.com"),
        ])
        self.assert_rejected(response, "البريد الإلكتروني", "مكرر")

    def test_email_of_another_user(self):
        User.objects.create_user(username="someone", email="taken@example.com")
        response = self.upload([student_row(email="taken@example.com")])
        self.assert_rejected(response, "البريد الإلكتروني", "مستخدم لحساب آخر")

    def test_unknown_department(self):
        response = self.upload([student_row(department="الهندسة")])
        self.assert_rejected(response, "القسم", "غير موجود")

    def test_unknown_specialization(self):
        response = self.upload([student_row(specialization="تصميم")])
        self.assert_rejected(response, "التخصص", "غير موجود")

    def test_specialization_from_another_department(self):
        response = self.upload([student_row(specialization="محاسبة")])
        self.assert_rejected(response, "التخصص", "لا يتبع")

    def test_negative_hours(self):
        response = self.upload([student_row(completed_credit_hours=-5)])
        self.assert_rejected(response, "الساعات المنجزة", "0 أو أكثر")

    def test_gpa_out_of_range(self):
        response = self.upload([student_row(gpa=5)])
        self.assert_rejected(response, "المعدل", "بين 0 و 4")

    def test_missing_required_value(self):
        response = self.upload([student_row(first_name=None)])
        self.assert_rejected(response, "الاسم الأول", "مطلوب")

    def test_missing_column(self):
        sheet = load_workbook(BytesIO(build_workbook([student_row()])))
        sheet.active.delete_cols(4)  # The email column.
        buffer = BytesIO()
        sheet.save(buffer)
        response = self.client.post(
            URL, {"file": SimpleUploadedFile("s.xlsx", buffer.getvalue())}, format="multipart"
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("البريد الإلكتروني", response.data["errors"][0]["message"])

    def test_only_xlsx_files(self):
        file = SimpleUploadedFile("students.csv", b"a,b,c")
        response = self.client.post(URL, {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("file", response.data)

    def test_damaged_xlsx_file(self):
        file = SimpleUploadedFile("students.xlsx", b"not really excel")
        response = self.client.post(URL, {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("تعذّرت قراءة الملف", response.data["errors"][0]["message"])

    def test_file_over_5mb(self):
        file = SimpleUploadedFile("students.xlsx", b"0" * (5 * 1024 * 1024 + 1))
        response = self.client.post(URL, {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("5", str(response.data["file"]))

    def test_existing_student_is_updated_but_keeps_password(self):
        user = User.objects.create_user(
            username="120260001", password="MyOwnPass1!", email="old@example.com", role=User.Role.STUDENT
        )
        StudentProfile.objects.create(user=user, completed_credit_hours=50)

        response = self.upload([student_row("120260001", email="new@example.com", completed_credit_hours=100)])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual((response.data["created"], response.data["updated"]), (0, 1))

        user.refresh_from_db()
        self.assertTrue(user.check_password("MyOwnPass1!"))
        self.assertEqual(user.email, "new@example.com")
        self.assertEqual(user.student_profile.completed_credit_hours, 100)
        self.assertEqual(User.objects.filter(username="120260001").count(), 1)

    def test_existing_user_with_another_role_is_not_changed(self):
        User.objects.create_user(username="120260001", role=User.Role.SUPERVISOR)
        response = self.upload([student_row("120260001")])
        self.assert_rejected(response, "الرقم الجامعي", "ليس طالبًا")
        self.assertEqual(User.objects.get(username="120260001").role, User.Role.SUPERVISOR)
        self.assertFalse(StudentProfile.objects.exists())

    def test_admin_can_import(self):
        admin = User.objects.create_user(username="admin", role=User.Role.ADMIN)
        self.client.force_authenticate(admin)
        self.assertEqual(self.upload([student_row()]).status_code, 200)

    def test_other_roles_are_forbidden(self):
        for role in (User.Role.STUDENT, User.Role.PROVIDER, User.Role.SUPERVISOR):
            self.client.force_authenticate(User.objects.create_user(username=role, role=role))
            self.assertEqual(self.upload([student_row()]).status_code, 403, role)
            self.assertEqual(self.client.get(TEMPLATE_URL).status_code, 403, role)
        self.assertFalse(User.objects.filter(username="120260001").exists())

    def test_must_be_logged_in(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(TEMPLATE_URL).status_code, 401)

    def test_template_downloads_right_to_left(self):
        response = self.client.get(TEMPLATE_URL)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual([c.value for c in sheet[1]], [header for header, _, _ in COLUMNS])

    def test_template_can_be_imported_as_is(self):
        response = self.client.get(TEMPLATE_URL)
        file = SimpleUploadedFile("template.xlsx", response.content)
        response = self.client.post(URL + "?dry_run=1", {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)

    def test_wafaa_student_urls_still_work(self):
        self.assertEqual(self.client.get("/api/students/").status_code, 200)
        student = User.objects.create_user(username="student", role=User.Role.STUDENT)
        self.client.force_authenticate(student)
        self.assertEqual(self.client.get("/api/students/me/").status_code, 200)


MEDIA = tempfile.mkdtemp()  # the demo reports' files


@override_settings(DEBUG=True, MEDIA_ROOT=MEDIA)
class SeedDemoTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def seed(self):
        call_command("seed_demo", stdout=StringIO())

    def counts(self):
        return [
            model.objects.count()
            for model in (
                Department, Specialization, Provider, Opportunity, User, StudentProfile,
                Application, Placement, ExceptionRequest, DailyLog, Report, Certificate,
            )
        ]

    def test_seed_creates_the_demo_data(self):
        self.seed()
        self.assertEqual(self.counts()[:6], [2, 4, 4, 6, 18, 14])
        self.assertEqual(Provider.objects.filter(status=Provider.Status.SUSPENDED).count(), 1)
        self.assertEqual(Opportunity.objects.filter(status=Opportunity.Status.CLOSED).count(), 1)
        eligible = [p for p in StudentProfile.objects.all() if check_eligibility(p)[0]]
        self.assertEqual(len(eligible), 11)
        provider_user = User.objects.get(username="demo_provider")
        self.assertEqual(provider_user.provider_profile.name, "شركة التقنية المتقدمة")
        for role in User.Role.values:
            self.assertTrue(User.objects.filter(role=role).exists(), role)

    def test_one_student_at_each_stage(self):
        self.seed()
        stages = dict(Application.objects.values_list("student__username", "status"))
        self.assertEqual(stages, {
            "120260001": Status.DRAFT,
            "120260002": Status.SUBMITTED,
            "120260003": Status.UNDER_REVIEW,
            "120260004": Status.ACCEPTED,
            "120260005": Status.REJECTED,
            "120260009": Status.ASSIGNED,
            "120260010": Status.IN_TRAINING,
            "120260011": Status.IN_TRAINING,
            "120260012": Status.COMPLETED,
            "120260013": Status.CERTIFIED,
            "120260014": Status.WITHDRAWN,
        })
        exceptions = dict(ExceptionRequest.objects.values_list("kind", "status"))
        self.assertEqual(exceptions, {
            "issue": "pending",
            "change_provider": "supervisor_approved",
            "withdrawal": "approved",
        })
        self.assertEqual(Certificate.objects.get().application.student.username, "120260013")
        # The demo accounts have something in their notifications.
        for username in ("demo_coordinator", "demo_supervisor", "demo_provider", "120260013"):
            self.assertTrue(Notification.objects.filter(recipient__username=username).exists(), username)

    def test_running_twice_adds_nothing(self):
        self.seed()
        first = self.counts()
        self.seed()
        self.assertEqual(self.counts(), first)

    @override_settings(DEBUG=False)
    def test_refuses_without_debug(self):
        with self.assertRaises(CommandError):
            self.seed()
        self.assertEqual(User.objects.count(), 0)

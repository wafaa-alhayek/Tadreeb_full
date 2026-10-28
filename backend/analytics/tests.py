from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.models import Department, Specialization, StudentProfile, User
from applications.models import Application, Placement
from applications.status import Status
from certificates.models import Certificate
from providers.models import Opportunity, Provider
from training.models import DailyLog


class AnalyticsTestData(APITestCase):
    def setUp(self):
        self.coordinator = User.objects.create_user(username="coord", password="x", role=User.Role.COORDINATOR)
        self.admin = User.objects.create_user(username="admin1", password="x", role=User.Role.ADMIN)
        self.sup = User.objects.create_user(username="sup1", password="x", role=User.Role.SUPERVISOR)
        self.it = Department.objects.create(name="تقنية المعلومات")
        self.business = Department.objects.create(name="إدارة الأعمال")
        specialization = Specialization.objects.create(department=self.it, name="برمجة")
        self.provider = Provider.objects.create(
            name="شركة التقنية", sector="تقنية", city="غزة", contact_name="م", contact_phone="0599000000"
        )
        self.opportunity = Opportunity.objects.create(
            provider=self.provider,
            specialization=specialization,
            title="تطوير الويب",
            tasks="مهام",
            seats=5,
            start_date=date(2026, 9, 1),
            end_date=date(2026, 12, 31),
            location="غزة",
        )
        self.training = self.make_application("20250001", "سارة", self.it, Status.IN_TRAINING)
        self.certified = self.make_application("20250002", "عمر", self.it, Status.CERTIFIED)
        self.rejected = self.make_application("20250003", "ليلى", self.business, Status.REJECTED)
        Placement.objects.create(
            application=self.training,
            opportunity=self.opportunity,
            start_date=self.opportunity.start_date,
            end_date=self.opportunity.end_date,
        )
        Placement.objects.create(
            application=self.certified,
            opportunity=self.opportunity,
            start_date=self.opportunity.start_date,
            end_date=self.opportunity.end_date,
            is_active=False,
        )
        DailyLog.objects.create(
            application=self.training, date=date(2026, 9, 2), hours=6, tasks="x", confirmed_at="2026-09-02T12:00Z"
        )
        DailyLog.objects.create(application=self.training, date=date(2026, 9, 3), hours=5, tasks="x")
        Certificate.objects.create(
            application=self.certified,
            serial_number="TDR-2026-00001",
            student_name="عمر",
            university_id="20250002",
            provider_name=self.provider.name,
            opportunity_title=self.opportunity.title,
            start_date=self.opportunity.start_date,
            end_date=self.opportunity.end_date,
        )
        self.client.force_authenticate(self.coordinator)

    def make_application(self, university_id, name, department, status):
        student = User.objects.create_user(
            username=university_id, password="x", role=User.Role.STUDENT, first_name=name
        )
        StudentProfile.objects.create(user=student, department=department)
        return Application.objects.create(student=student, status=status)


class StatsTests(AnalyticsTestData):
    def test_dashboard_numbers(self):
        response = self.client.get("/api/stats/")
        self.assertEqual(response.status_code, 200)
        data = response.data
        self.assertEqual(data["students"], 3)
        self.assertEqual(data["applications"]["total"], 3)
        self.assertEqual(data["applications"]["by_status"][Status.IN_TRAINING], 1)
        self.assertEqual(data["applications"]["by_status"][Status.DRAFT], 0)
        self.assertEqual(data["opportunities"], {"open": 1, "seats": 5, "free_seats": 4})
        self.assertEqual(data["placements"]["active"], 1)
        self.assertEqual(data["training"]["confirmed_hours"], 6)
        self.assertEqual(data["training"]["pending_days"], 1)
        self.assertEqual(data["certificates"], {"issued": 1, "revoked": 0})
        departments = {row["department"]: row for row in data["by_department"]}
        self.assertEqual(departments["تقنية المعلومات"]["applications"], 2)
        self.assertEqual(departments["تقنية المعلومات"]["certified"], 1)
        self.assertEqual(data["by_provider"], [{"id": self.provider.id, "name": "شركة التقنية", "active_placements": 1}])

    def test_filter_by_department(self):
        response = self.client.get(f"/api/stats/?department={self.business.id}")
        self.assertEqual(response.data["applications"]["total"], 1)
        self.assertEqual(response.data["certificates"]["issued"], 0)

    def test_only_coordinator_and_admin(self):
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.get("/api/stats/").status_code, 200)
        for user in (self.sup, self.training.student):
            self.client.force_authenticate(user)
            self.assertEqual(self.client.get("/api/stats/").status_code, 403)


class ExportTests(AnalyticsTestData):
    def read(self, response):
        sheet = load_workbook(BytesIO(response.content)).active
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    def test_export_applications_to_excel(self):
        response = self.client.get("/api/exports/applications/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        rows = self.read(response)
        self.assertEqual(rows[0][:3], ["رقم الطلب", "الرقم الجامعي", "اسم الطالب"])
        self.assertEqual(len(rows), 4)
        by_id = {row[1]: row for row in rows[1:]}
        self.assertEqual(by_id["20250001"][6], "شركة التقنية")
        self.assertEqual(by_id["20250001"][11], 6)
        self.assertEqual(by_id["20250002"][13], "TDR-2026-00001")

    def test_export_uses_the_filters(self):
        rows = self.read(self.client.get(f"/api/exports/applications/?status={Status.REJECTED}"))
        self.assertEqual([row[1] for row in rows[1:]], ["20250003"])

    def test_supervisor_cannot_export(self):
        self.client.force_authenticate(self.sup)
        self.assertEqual(self.client.get("/api/exports/applications/").status_code, 403)


class ArchiveTests(AnalyticsTestData):
    def ids(self, query=""):
        response = self.client.get(f"/api/archive/{query}")
        self.assertEqual(response.status_code, 200)
        return {row["id"] for row in response.data["results"]}

    def test_only_finished_files(self):
        self.assertEqual(self.ids(), {self.certified.id, self.rejected.id})

    def test_search_and_filters(self):
        self.assertEqual(self.ids("?search=20250002"), {self.certified.id})
        self.assertEqual(self.ids("?search=ليلى"), {self.rejected.id})
        self.assertEqual(self.ids(f"?provider={self.provider.id}"), {self.certified.id})
        self.assertEqual(self.ids(f"?department={self.business.id}"), {self.rejected.id})
        self.assertEqual(self.ids("?year=1999"), set())

    def test_archive_row(self):
        response = self.client.get("/api/archive/?search=20250002")
        row = response.data["results"][0]
        self.assertEqual(row["provider_name"], "شركة التقنية")
        self.assertEqual(row["certificate_serial"], "TDR-2026-00001")
        self.assertEqual(row["department"], "تقنية المعلومات")
